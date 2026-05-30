from functools import wraps
from flask import render_template, redirect, url_for, flash, request, abort, send_file
from flask_login import login_required, current_user
from app import db
from app.admin import bp
from app.models import Oferta, Contacto, Queja, Usuario, HistorialOferta, Reserva
from app.forms import OfertaForm
from app.services.file_service import save_uploaded_files, delete_files
from app.ciudades import (
    CIUDADES_GUINEA_ECUATORIAL, TIPOS_PROPIEDAD, SUBTIPOS_VIVIENDA,
    TIPOLOGIAS_LOCAL, TIPOLOGIAS_OFICINA, TIPOS_ALQUILER, ESTADOS_PROPIEDAD
)
from datetime import datetime, date
import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def admin_required(f):
    """Decorator to require admin role"""
    @wraps(f)
    @login_required
    def decorated_function(*args, **kwargs):
        if not current_user.is_admin():
            abort(403)
        return f(*args, **kwargs)
    return decorated_function


@bp.route('/dashboard')
@admin_required
def dashboard():
    """Admin dashboard with statistics"""
    # Statistics
    total_ofertas = Oferta.query.count()
    ofertas_activas = Oferta.query.filter_by(estado='activo').count()
    total_contactos = Contacto.query.count()
    contactos_pendientes = Contacto.query.filter_by(estado='pendiente').count()
    total_quejas = Queja.query.count()
    quejas_pendientes = Queja.query.filter_by(estado='pendiente').count()
    total_usuarios = Usuario.query.count()
    
    # Recent contacts
    contactos_recientes = Contacto.query.order_by(Contacto.fecha_contacto.desc()).limit(10).all()
    
    # Recent properties
    ofertas_recientes = Oferta.query.order_by(Oferta.fecha_creacion.desc()).limit(5).all()
    
    return render_template('admin/dashboard.html',
                         title='Panel de Administración',
                         total_ofertas=total_ofertas,
                         ofertas_activas=ofertas_activas,
                         total_contactos=total_contactos,
                         contactos_pendientes=contactos_pendientes,
                         total_quejas=total_quejas,
                         quejas_pendientes=quejas_pendientes,
                         total_usuarios=total_usuarios,
                         contactos_recientes=contactos_recientes,
                         ofertas_recientes=ofertas_recientes)


@bp.route('/ofertas')
@admin_required
def ofertas():
    """Manage all property listings"""
    page = request.args.get('page', 1, type=int)
    
    query = Oferta.query
    
    # Filter by search query (Title, Description, or Location)
    q = request.args.get('q', '')
    if q:
        query = query.filter(db.or_(
            Oferta.titulo.ilike(f'%{q}%'),
            Oferta.descripcion.ilike(f'%{q}%'),
            Oferta.ciudad.ilike(f'%{q}%'),
            Oferta.barrio.ilike(f'%{q}%')
        ))
    
    # Filter by status
    estado = request.args.get('estado', '')
    if estado:
        query = query.filter_by(estado=estado)
        
    # Filter by transaction type
    tipo_operacion = request.args.get('tipo_operacion', '')
    if tipo_operacion:
        query = query.filter_by(tipo_operacion=tipo_operacion)
        
    # Filter by property type
    tipo_propiedad = request.args.get('tipo_propiedad', '')
    if tipo_propiedad:
        query = query.filter_by(tipo_propiedad=tipo_propiedad)
    
    pagination = query.order_by(Oferta.fecha_creacion.desc()).paginate(
        page=page, per_page=20, error_out=False
    )
    
    return render_template('admin/ofertas.html',
                         title='Gestión de Ofertas',
                         ofertas=pagination.items,
                         pagination=pagination)


@bp.route('/ofertas/bulk', methods=['POST'])
@admin_required
def ofertas_bulk():
    """Bulk actions on property listings"""
    action = request.form.get('action')
    ids = request.form.getlist('ids')
    
    if not ids:
        flash('No se seleccionaron ofertas.', 'warning')
        return redirect(url_for('admin.ofertas'))
        
    ofertas_sel = Oferta.query.filter(Oferta.id.in_(ids)).all()
    
    if action == 'activar':
        for o in ofertas_sel:
            o.estado = 'activada'
        flash(f'Se han activado {len(ofertas_sel)} ofertas.', 'success')
    elif action == 'desactivar':
        for o in ofertas_sel:
            o.estado = 'desactivada'
        flash(f'Se han desactivado {len(ofertas_sel)} ofertas.', 'success')
    elif action == 'borrar':
        count = 0
        for o in ofertas_sel:
            # Delete associated files
            delete_files(o.get_imagenes())
            delete_files(o.get_videos())
            db.session.delete(o)
            count += 1
        flash(f'Se han eliminado {count} ofertas.', 'success')
    
    db.session.commit()
    return redirect(url_for('admin.ofertas'))


@bp.route('/ofertas/nueva', methods=['GET', 'POST'])
@admin_required
def oferta_nueva():
    """Create new property listing"""
    form = OfertaForm()
    
    if form.validate_on_submit():
        # Create new oferta
        oferta = Oferta(
            titulo=form.titulo.data,
            descripcion=form.descripcion.data,
            precio=form.precio.data,
            ubicacion=form.ubicacion.data,
            latitud=form.latitud.data,
            longitud=form.longitud.data,
            tipo_propiedad=form.tipo_propiedad.data,
            subtipo_propiedad=form.subtipo_propiedad.data,
            tipo_operacion=form.tipo_operacion.data,
            tipo_alquiler=form.tipo_alquiler.data,
            estado_propiedad=form.estado_propiedad.data,
            metros_cuadrados=form.m2.data,
            estado=form.estado.data,
            destacado=form.destacado.data,
            fecha_baja_programada=form.fecha_baja_programada.data,
            propietario_id=current_user.id
        )
        
        # Mapping ubicacion (from SelectField) to ciudad column
        oferta.ciudad = dict(get_ciudades_choices()).get(form.ubicacion.data, form.ubicacion.data)
        oferta.barrio = form.barrio.data
        
        # Set characteristics
        caracteristicas = {
            'habitaciones': form.habitaciones.data or '0',
            'banos': form.banos.data or '0',
            'm2': form.m2.data or 0,
            'vivienda_items': form.caracteristicas_vivienda.data or [],
            'oficina_items': form.caracteristicas_oficina.data or [],
            'local_items': form.caracteristicas_local.data or [],
            'terreno_items': form.caracteristicas_terreno.data or []
        }
        oferta.set_caracteristicas(caracteristicas)
        
        # Handle file uploads
        imagenes = save_uploaded_files(form.imagenes.data, 'images')
        videos = save_uploaded_files(form.videos.data, 'videos')
        oferta.set_imagenes(imagenes)
        oferta.set_videos(videos)
        
        db.session.add(oferta)
        db.session.commit()
        
        flash('Oferta creada exitosamente.', 'success')
        return redirect(url_for('admin.ofertas'))
    
    return render_template('admin/oferta_form.html',
                         title='Nueva Oferta',
                         form=form,
                         accion='Crear',
                         now=datetime.now(),
                         HistorialOferta=HistorialOferta)


@bp.route('/ofertas/editar/<int:id>', methods=['GET', 'POST'])
@admin_required
def oferta_editar(id):
    """Edit existing property listing"""
    oferta = Oferta.query.get_or_404(id)
    form = OfertaForm(obj=oferta)
    
    if request.method == 'GET':
        # Pre-fill characteristics
        caract = oferta.get_caracteristicas()
        form.habitaciones.data = caract.get('habitaciones', '')
        form.banos.data = caract.get('banos', '')
        form.m2.data = caract.get('m2', 0)
        form.caracteristicas_vivienda.data = caract.get('vivienda_items', [])
        form.caracteristicas_oficina.data = caract.get('oficina_items', [])
        form.caracteristicas_local.data = caract.get('local_items', [])
        form.caracteristicas_terreno.data = caract.get('terreno_items', [])
        form.barrio.data = oferta.barrio
    
    if form.validate_on_submit():
        # Track changes for history
        cambios = {}
        
        if oferta.titulo != form.titulo.data:
            cambios['titulo'] = (oferta.titulo, form.titulo.data)
        if oferta.precio != form.precio.data:
            cambios['precio'] = (oferta.precio, form.precio.data)
        
        # Update oferta
        oferta.titulo = form.titulo.data
        oferta.descripcion = form.descripcion.data
        oferta.precio = form.precio.data
        oferta.ubicacion = form.ubicacion.data
        oferta.latitud = form.latitud.data
        oferta.longitud = form.longitud.data
        oferta.tipo_propiedad = form.tipo_propiedad.data
        oferta.subtipo_propiedad = form.subtipo_propiedad.data
        oferta.tipo_operacion = form.tipo_operacion.data
        oferta.tipo_alquiler = form.tipo_alquiler.data
        oferta.estado_propiedad = form.estado_propiedad.data
        oferta.metros_cuadrados = form.m2.data
        oferta.ciudad = form.ubicacion.data # We use the selection directly for now
        oferta.barrio = form.barrio.data
        oferta.estado = form.estado.data
        oferta.destacado = form.destacado.data
        oferta.fecha_baja_programada = form.fecha_baja_programada.data
        
        # Update characteristics
        caracteristicas = {
            'habitaciones': form.habitaciones.data or '0',
            'banos': form.banos.data or '0',
            'm2': form.m2.data or 0,
            'vivienda_items': form.caracteristicas_vivienda.data or [],
            'oficina_items': form.caracteristicas_oficina.data or [],
            'local_items': form.caracteristicas_local.data or [],
            'terreno_items': form.caracteristicas_terreno.data or []
        }
        oferta.set_caracteristicas(caracteristicas)
        
        # Handle new file uploads
        if form.imagenes.data and form.imagenes.data[0].filename:
            nuevas_imagenes = save_uploaded_files(form.imagenes.data, 'images')
            imagenes_actuales = oferta.get_imagenes()
            imagenes_actuales.extend(nuevas_imagenes)
            oferta.set_imagenes(imagenes_actuales)
        
        if form.videos.data and form.videos.data[0].filename:
            nuevos_videos = save_uploaded_files(form.videos.data, 'videos')
            videos_actuales = oferta.get_videos()
            videos_actuales.extend(nuevos_videos)
            oferta.set_videos(videos_actuales)
        
        # Record changes in history
        for campo, (anterior, nuevo) in cambios.items():
            historial = HistorialOferta(
                oferta_id=oferta.id,
                campo_modificado=campo,
                valor_anterior=str(anterior),
                valor_nuevo=str(nuevo),
                usuario_modificador_id=current_user.id
            )
            db.session.add(historial)
        
        db.session.commit()
        
        flash('Oferta actualizada exitosamente.', 'success')
        return redirect(url_for('admin.ofertas'))
    
    return render_template('admin/oferta_form.html',
                         title='Editar Oferta',
                         form=form,
                         oferta=oferta,
                         accion='Editar',
                         now=datetime.now(),
                         HistorialOferta=HistorialOferta)


@bp.route('/ofertas/eliminar/<int:id>', methods=['POST'])
@admin_required
def oferta_eliminar(id):
    """Delete property listing"""
    oferta = Oferta.query.get_or_404(id)
    
    # Delete associated files
    delete_files(oferta.get_imagenes())
    delete_files(oferta.get_videos())
    
    db.session.delete(oferta)
    db.session.commit()
    
    flash('Oferta eliminada exitosamente.', 'success')
    return redirect(url_for('admin.ofertas'))


@bp.route('/ofertas/exportar')
@admin_required
def exportar_ofertas_csv():
    """Export property listings to CSV"""
    ofertas_to_export = Oferta.query.order_by(Oferta.fecha_creacion.desc()).all()
    
    # Create CSV in memory
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header
    writer.writerow([
        'ID', 'Título', 'Ubicación', 'Barrio', 'Tipo Transacción', 'Tipo Propiedad', 
        'Estado de la Oferta', 'Precio (FCFA)', 'Tamaño (m2)', 'Fecha de Alta', 
        'Baja Programada', 'Última Modificación', 'Usuario'
    ])
    
    # Data
    for o in ofertas_to_export:
        # Get creator name
        usuario_name = o.propietario.nombre if o.propietario else 'Sistema'
        
        # Check for last modifier in history
        last_mod = HistorialOferta.query.filter_by(oferta_id=o.id).order_by(HistorialOferta.fecha_modificacion.desc()).first()
        if last_mod and last_mod.usuario_modificador:
            usuario_name = last_mod.usuario_modificador.nombre
            
        writer.writerow([
            o.id,
            o.titulo,
            o.ciudad,
            o.barrio,
            o.tipo_operacion,
            o.tipo_propiedad,
            o.estado,
            o.precio,
            o.metros_cuadrados,
            o.fecha_creacion.strftime('%Y-%m-%d %H:%M'),
            o.fecha_baja_programada.strftime('%Y-%m-%d') if o.fecha_baja_programada else '',
            o.fecha_modificacion.strftime('%Y-%m-%d %H:%M'),
            usuario_name
        ])
    
    # Prepare file for download
    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'ofertas_segis_{datetime.now().strftime("%Y%m%d")}.csv'
    )


@bp.route('/ofertas/plantilla-excel')
@admin_required
def plantilla_excel():
    """Download a pre-filled Excel template for bulk offer import"""
    wb = openpyxl.Workbook()

    # ---- Sheet 1: Data entry ----
    ws = wb.active
    ws.title = 'Ofertas'

    header_fill = PatternFill('solid', fgColor='1565C0')
    header_font = Font(bold=True, color='FFFFFF', size=10)
    note_fill  = PatternFill('solid', fgColor='E3F2FD')
    thin = Side(style='thin', color='BDBDBD')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    columns = [
        ('titulo',            'Título (*)',                 40),
        ('descripcion',       'Descripción (*)',            50),
        ('precio',            'Precio FCFA (*)',            16),
        ('ciudad',            'Ciudad (*)',                 20),
        ('barrio',            'Barrio/Zona',               20),
        ('tipo_operacion',    'Tipo Transacción (*)',       20),
        ('tipo_propiedad',    'Tipo Propiedad (*)',         20),
        ('subtipo_propiedad', 'Subtipo Propiedad',         22),
        ('estado_propiedad',  'Estado Propiedad',          20),
        ('tipo_alquiler',     'Tipo Alquiler',             18),
        ('m2',                'Tamaño m² (*)',             14),
        ('habitaciones',      'N.º Habitaciones',          18),
        ('banos',             'N.º Baños',                 14),
        ('caracteristicas',   'Características (coma)',    40),
        ('latitud',           'Latitud GPS',               14),
        ('longitud',          'Longitud GPS',              14),
        ('fecha_baja',        'Fecha Baja (YYYY-MM-DD)',   22),
        ('destacado',         'Destacado (si/no)',         18),
    ]

    for col_idx, (field, header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill   = header_fill
        cell.font   = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30

    # Example row
    example = [
        'Apartamento céntrico luminoso',
        'Amplio apartamento de 3 habitaciones en zona residencial con vistas al mar.',
        500000,
        'Malabo',
        'Centro',
        'alquiler',
        'vivienda',
        'apartamento',
        'muy_bien',
        'larga_duracion',
        85,
        '3',
        '2',
        'amueblado,parking,internet',
        3.7500,
        8.7833,
        '',
        'no',
    ]
    for col_idx, val in enumerate(example, start=1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.fill   = note_fill
        cell.border = border
        cell.alignment = Alignment(wrap_text=True)

    ws.freeze_panes = 'A2'

    # ---- Sheet 2: Valid values reference ----
    ws2 = wb.create_sheet('Valores Válidos')
    ws2['A1'] = 'Campo'
    ws2['B1'] = 'Valores válidos'
    for cell in [ws2['A1'], ws2['B1']]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2E7D32')

    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 60

    ref_rows = [
        ('ciudad',           ', '.join(c[0] for c in CIUDADES_GUINEA_ECUATORIAL)),
        ('tipo_operacion',   'venta, alquiler, compartir, booking'),
        ('tipo_propiedad',   ', '.join(v[0] for v in TIPOS_PROPIEDAD if v[0])),
        ('subtipo_propiedad',', '.join(v[0] for v in SUBTIPOS_VIVIENDA + TIPOLOGIAS_LOCAL + TIPOLOGIAS_OFICINA if v[0])),
        ('estado_propiedad', ', '.join(v[0] for v in ESTADOS_PROPIEDAD if v[0])),
        ('tipo_alquiler',    ', '.join(v[0] for v in TIPOS_ALQUILER if v[0])),
        ('habitaciones',     '0, 1, 2, 3, 4+'),
        ('banos',            '0, 1, 2, 3, 4+'),
        ('caracteristicas',  'amueblado, sin_amueblar, con_electrodomesticos, sin_electrodomesticos, accesibilidad, '
                             'ascensor, internet, aire_acondicionado, jardin, parking, terraza, piscina, '
                             'acceso_rodado, sin_acceso_rodado, terreno_llano, terreno_montanoso, '
                             'suministro_electrico, cerca_centro'),
        ('destacado',        'si, no'),
    ]
    for row_idx, (campo, vals) in enumerate(ref_rows, start=2):
        ws2.cell(row=row_idx, column=1, value=campo)
        cell = ws2.cell(row=row_idx, column=2, value=vals)
        cell.alignment = Alignment(wrap_text=True)
        ws2.row_dimensions[row_idx].height = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='plantilla_importacion_ofertas.xlsx'
    )


@bp.route('/ofertas/importar-excel', methods=['POST'])
@admin_required
def importar_excel():
    """Process uploaded Excel file and bulk-create property listings"""
    archivo = request.files.get('excel_file')
    if not archivo or not archivo.filename.endswith('.xlsx'):
        flash('Por favor sube un archivo .xlsx válido.', 'danger')
        return redirect(url_for('admin.ofertas'))

    try:
        wb = openpyxl.load_workbook(archivo, data_only=True)
        ws = wb.active
    except Exception as e:
        flash(f'No se pudo leer el archivo Excel: {e}', 'danger')
        return redirect(url_for('admin.ofertas'))

    VALID_OPERACIONES = {'venta', 'alquiler', 'compartir', 'booking'}
    VALID_PROPIEDADES = {v[0] for v in TIPOS_PROPIEDAD if v[0]}
    CIUDADES_SET      = {c[0] for c in CIUDADES_GUINEA_ECUATORIAL}

    created = 0
    errors  = []

    # Skip header row (row 1), process from row 2
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        # Skip completely empty rows
        if all(v is None or str(v).strip() == '' for v in row):
            continue

        (
            titulo, descripcion, precio, ciudad, barrio,
            tipo_operacion, tipo_propiedad, subtipo_propiedad,
            estado_propiedad, tipo_alquiler, m2, habitaciones,
            banos, caracteristicas_raw, latitud, longitud,
            fecha_baja_raw, destacado_raw
        ) = (list(row) + [None] * 18)[:18]

        row_errors = []

        # --- Validation ---
        titulo = str(titulo).strip() if titulo else ''
        if not titulo:
            row_errors.append('título vacío')

        descripcion = str(descripcion).strip() if descripcion else ''

        try:
            precio = float(precio)
            if precio <= 0:
                raise ValueError
        except (TypeError, ValueError):
            row_errors.append('precio inválido (debe ser número > 0)')
            precio = 0

        ciudad = str(ciudad).strip() if ciudad else ''
        if not ciudad:
            row_errors.append('ciudad vacía')

        tipo_operacion = str(tipo_operacion).strip().lower() if tipo_operacion else ''
        if tipo_operacion not in VALID_OPERACIONES:
            row_errors.append(f'tipo_operacion inválido: "{tipo_operacion}" (usa: venta, alquiler, compartir, booking)')

        tipo_propiedad = str(tipo_propiedad).strip().lower() if tipo_propiedad else ''
        if tipo_propiedad not in VALID_PROPIEDADES:
            row_errors.append(f'tipo_propiedad inválido: "{tipo_propiedad}"')

        try:
            m2 = float(m2) if m2 else 0
        except (TypeError, ValueError):
            m2 = 0

        if row_errors:
            errors.append(f'Fila {row_num}: ' + '; '.join(row_errors))
            continue

        # --- Parse optional fields ---
        subtipo_propiedad  = str(subtipo_propiedad).strip()  if subtipo_propiedad  else ''
        estado_propiedad   = str(estado_propiedad).strip()   if estado_propiedad   else ''
        tipo_alquiler_val  = str(tipo_alquiler).strip()      if tipo_alquiler      else ''
        barrio_val         = str(barrio).strip()             if barrio             else ''
        habitaciones_val   = str(habitaciones).strip()       if habitaciones       else ''
        banos_val          = str(banos).strip()              if banos              else ''
        destacado_val      = str(destacado_raw).strip().lower() == 'si' if destacado_raw else False

        try:
            lat = float(latitud)  if latitud  else None
            lng = float(longitud) if longitud else None
        except (TypeError, ValueError):
            lat = lng = None

        fecha_baja = None
        if fecha_baja_raw:
            try:
                if isinstance(fecha_baja_raw, (datetime, date)):
                    fecha_baja = fecha_baja_raw if isinstance(fecha_baja_raw, date) else fecha_baja_raw.date()
                else:
                    fecha_baja = datetime.strptime(str(fecha_baja_raw).strip(), '%Y-%m-%d').date()
            except ValueError:
                pass  # Invalid date — ignore

        # Parse comma-separated characteristics
        caract_items = [c.strip() for c in str(caracteristicas_raw).split(',') if c.strip()] \
            if caracteristicas_raw else []
        caracteristicas = {
            'habitaciones': habitaciones_val or '0',
            'banos': banos_val or '0',
            'm2': m2,
            'vivienda_items': caract_items if tipo_propiedad == 'vivienda' else [],
            'oficina_items':  caract_items if tipo_propiedad == 'oficinas' else [],
            'local_items':    caract_items if tipo_propiedad == 'locales_naves' else [],
            'terreno_items':  caract_items if tipo_propiedad == 'terrenos' else [],
        }

        oferta = Oferta(
            titulo=titulo,
            descripcion=descripcion,
            precio=precio,
            ubicacion=ciudad,
            ciudad=ciudad,
            barrio=barrio_val,
            latitud=lat,
            longitud=lng,
            tipo_propiedad=tipo_propiedad,
            subtipo_propiedad=subtipo_propiedad,
            tipo_operacion=tipo_operacion,
            tipo_alquiler=tipo_alquiler_val,
            estado_propiedad=estado_propiedad,
            metros_cuadrados=m2,
            estado='abierta',
            destacado=destacado_val,
            fecha_baja_programada=fecha_baja,
            propietario_id=current_user.id
        )
        oferta.set_caracteristicas(caracteristicas)
        oferta.set_imagenes([])
        oferta.set_videos([])
        db.session.add(oferta)
        created += 1

    if created:
        db.session.commit()
        flash(f'✅ {created} oferta(s) importada(s) correctamente. Recuerda añadir las imágenes a cada oferta.', 'success')
    if errors:
        for err in errors:
            flash(f'⚠️ {err}', 'warning')
    if not created and not errors:
        flash('El archivo no contenía filas de datos.', 'info')

    return redirect(url_for('admin.ofertas'))


@bp.route('/buzon')
@admin_required
def buzon():
    """Contact inbox"""
    page = request.args.get('page', 1, type=int)
    
    query = Contacto.query
    
    # Filters
    estado = request.args.get('estado', '')
    if estado:
        query = query.filter_by(estado=estado)
    
    fecha_desde = request.args.get('fecha_desde', '')
    if fecha_desde:
        query = query.filter(Contacto.fecha_contacto >= datetime.strptime(fecha_desde, '%Y-%m-%d'))
    
    pagination = query.order_by(Contacto.fecha_contacto.desc()).paginate(
        page=page, per_page=20, error_out=False
    )
    
    return render_template('admin/buzon.html',
                         title='Buzón de Contactos',
                         contactos=pagination.items,
                         pagination=pagination)


@bp.route('/buzon/<int:id>/resolver', methods=['POST'])
@admin_required
def contacto_resolver(id):
    """Mark contact as resolved"""
    contacto = Contacto.query.get_or_404(id)
    contacto.estado = 'resuelto'
    contacto.fecha_resolucion = datetime.utcnow()
    contacto.notas_admin = request.form.get('notas', '')
    
    db.session.commit()
    
    flash('Contacto marcado como resuelto.', 'success')
    return redirect(url_for('admin.buzon'))


@bp.route('/buzon/exportar')
@admin_required
def exportar_contactos():
    """Export contacts to CSV"""
    contactos = Contacto.query.order_by(Contacto.fecha_contacto.desc()).all()
    
    # Create CSV in memory
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Header
    writer.writerow(['ID', 'Fecha', 'Nombre', 'Email', 'Teléfono', 'Propiedad', 'Estado', 'Mensaje'])
    
    # Data
    for c in contactos:
        writer.writerow([
            c.id,
            c.fecha_contacto.strftime('%Y-%m-%d %H:%M'),
            c.nombre,
            c.email,
            c.telefono,
            c.oferta.titulo,
            c.estado,
            c.mensaje[:100]  # Truncate message
        ])
    
    # Prepare file for download
    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'contactos_{datetime.now().strftime("%Y%m%d")}.csv'
    )


@bp.route('/quejas')
@admin_required
def quejas():
    """Complaint management"""
    page = request.args.get('page', 1, type=int)
    
    query = Queja.query
    
    # Filter by status
    estado = request.args.get('estado', '')
    if estado:
        query = query.filter_by(estado=estado)
    
    pagination = query.order_by(Queja.fecha_creacion.desc()).paginate(
        page=page, per_page=20, error_out=False
    )
    
    return render_template('admin/quejas.html',
                         title='Gestión de Quejas',
                         quejas=pagination.items,
                         pagination=pagination)


@bp.route('/quejas/<int:id>/responder', methods=['POST'])
@admin_required
def queja_responder(id):
    """Respond to complaint"""
    queja = Queja.query.get_or_404(id)
    queja.respuesta_admin = request.form.get('respuesta', '')
    queja.estado = 'resuelto'
    queja.fecha_resolucion = datetime.utcnow()
    
    db.session.commit()
    
    # TODO: Send email to user
    
    flash('Respuesta enviada correctamente.', 'success')
    return redirect(url_for('admin.quejas'))


@bp.route('/usuarios')
@admin_required
def usuarios():
    """View all registered users"""
    page = request.args.get('page', 1, type=int)
    
    query = Usuario.query
    
    # Filter by role
    rol = request.args.get('rol', '')
    if rol:
        query = query.filter_by(rol=rol)
    
    pagination = query.order_by(Usuario.fecha_registro.desc()).paginate(
        page=page, per_page=20, error_out=False
    )
    
    return render_template('admin/usuarios.html',
                         title='Gestión de Usuarios',
                         usuarios=pagination.items,
                         pagination=pagination)

@bp.route('/reservas')
@admin_required
def reservas():
    """Manage reservations"""
    page = request.args.get('page', 1, type=int)
    
    query = Reserva.query
    
    # Filter by status
    estado = request.args.get('estado', '')
    if estado:
        query = query.filter_by(estado=estado)
        
    pagination = query.order_by(Reserva.fecha_creacion.desc()).paginate(
        page=page, per_page=20, error_out=False
    )
    
    return render_template('admin/reservas.html',
                         title='Gestión de Reservas',
                         reservas=pagination.items,
                         pagination=pagination)

@bp.route('/reservas/<int:id>/actualizar', methods=['POST'])
@admin_required
def reserva_actualizar(id):
    """Update reservation status"""
    reserva = Reserva.query.get_or_404(id)
    nuevo_estado = request.form.get('estado')
    
    if nuevo_estado in ['pendiente', 'confirmada', 'cancelada', 'completada']:
        reserva.estado = nuevo_estado
        db.session.commit()
        flash(f'Estado de reserva actualizado a {nuevo_estado}.', 'success')
    else:
        flash('Estado no válido.', 'danger')
        
    return redirect(url_for('admin.reservas'))
